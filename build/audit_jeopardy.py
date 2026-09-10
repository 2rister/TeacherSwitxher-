"""Full audit of the Jeopardy deck: navigation + content.

Run it on the repaired deck (the default) or pass a path.
"""
import os, re, sys
from pptx import Presentation
from pptx.util import Emu

def audit(path, quiet=False):
    p = Presentation(path)
    S = p.slides
    idx = {sl.slide_id: i + 1 for i, sl in enumerate(S)}   # Slide is unhashable
    def head(i): 
        t = [sh.text_frame.text.strip() for sh in S[i-1].shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        return t[0] if t else ""
    def texts(i):
        return [sh.text_frame.text.strip() for sh in S[i-1].shapes
                if sh.has_text_frame and sh.text_frame.text.strip()]

    def target(sh):
        try:
            ts = sh.click_action.target_slide
            if ts is not None:
                return idx[ts.slide_id]
        except Exception:
            pass
        return None

    QS = [i for i in range(1, len(S) + 1) if "POINTS" in head(i)]
    AS = [i for i in range(1, len(S) + 1) if head(i).startswith("ANSWER")]
    problems = []

    # --- board grid ------------------------------------------------------
    CATS = ["CHILDHOOD & EARLY LIFE", "CAREER", "ACHIEVEMENTS", "CURIOUS FACTS", "WILD CARD"]
    tiles = []
    for sh in S[1].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().isdigit():
            tiles.append((round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2),
                          int(sh.text_frame.text.strip()), sh))
    cols = sorted({t[0] for t in tiles})
    rows = sorted({t[1] for t in tiles})
    if not quiet:
        print(f"board: {len(tiles)} tiles, {len(cols)} columns x {len(rows)} rows")
    for x, y, pts, sh in tiles:
        ci, ri = cols.index(x), rows.index(y)
        want_cat, want_pts = CATS[ci], (ri + 1) * 100
        if pts != want_pts:
            problems.append(f"board: tile at column {ci+1} row {ri+1} reads {pts}, grid says {want_pts}")
        dest = target(sh)
        if dest is None:
            problems.append(f"board: tile {want_cat} {want_pts} has no link"); continue
        h = head(dest)
        m = re.match(r"(.+?)\s*[·•]\s*(\d+)\s*POINTS", h)
        if not m:
            problems.append(f"board: tile {want_cat} {want_pts} -> slide {dest}, not a question")
        elif m.group(1).strip().upper() != want_cat or int(m.group(2)) != want_pts:
            problems.append(f"board: {want_cat} {want_pts} opens slide {dest} = "
                            f"{m.group(1).strip()} {m.group(2)}")

    # --- question / answer slides ---------------------------------------
    for i in QS:
        shapes = {sh.text_frame.text.strip().upper(): sh for sh in S[i-1].shapes
                  if sh.has_text_frame and sh.text_frame.text.strip()}
        show = next((sh for k, sh in shapes.items() if "SHOW" in k), None)
        back = next((sh for k, sh in shapes.items() if "BOARD" in k), None)
        if show is None: problems.append(f"slide {i}: no SHOW ANSWER button")
        elif target(show) != i + 1:
            problems.append(f"slide {i}: SHOW ANSWER -> {target(show)}, expected {i+1}")
        if back is None: problems.append(f"slide {i}: no BACK TO BOARD button")
        elif target(back) != 2:
            problems.append(f"slide {i}: BACK TO BOARD -> {target(back)}, expected 2")
        qm = re.match(r"(.+?)\s*[·•]\s*(\d+)\s*POINTS", head(i))
        am = re.match(r"ANSWER\s*[·•]\s*(.+?)\s*[·•]\s*(\d+)", head(i + 1)) if i + 1 in AS else None
        if qm and am and (qm.group(1).strip().lower() != am.group(1).strip().lower()
                          or qm.group(2) != am.group(2)):
            problems.append(f"slide {i}: answer slide {i+1} is labelled '{am.group(1)} {am.group(2)}'")

    for i in AS + [i for i in range(1, len(S)+1) if "FINAL JEOPARDY" in head(i)]:
        back = next((sh for sh in S[i-1].shapes
                     if sh.has_text_frame and "BOARD" in sh.text_frame.text.upper()), None)
        if back is None: problems.append(f"slide {i}: no BACK TO BOARD button")
        elif target(back) != 2:
            problems.append(f"slide {i}: BACK TO BOARD -> {target(back)}, expected 2")

    # --- content ---------------------------------------------------------
    qtexts = {}
    for i in QS:
        body = [t for t in texts(i) if t != head(i) and "SHOW" not in t.upper()
                and "BOARD" not in t.upper()]
        q = body[0] if body else ""
        key = re.sub(r"[^a-z0-9 ]", "", q.lower()).strip()
        if key in qtexts:
            problems.append(f"slide {i}: repeats the question on slide {qtexts[key]}")
        qtexts[key] = i
        # An identify-the-person question must not name the answer's gender:
        # with three men and two women on the board, "his" halves the field.
        if re.match(r"(who|which person)\b", q, re.I) and re.search(
                r"\b(his|her|hers|he|she|him)\b", q, re.I):
            problems.append(f"slide {i}: a pronoun gives the answer's gender away -> {q!r}")

    for i in AS:
        a = " ".join(texts(i))
        if "Newton" in a and "Nelson" in a and "Elizabeth" in a and "→" in a:
            order = re.findall(r"(Elizabeth|Newton|Nelson|Lennon|Diana)", a)
            correct = ["Elizabeth", "Newton", "Nelson", "Lennon", "Diana"]
            if order[:5] != correct:
                problems.append(f"slide {i}: birth order is {' -> '.join(order[:5])}, "
                                f"correct is {' -> '.join(correct)}")
    problems += xml_link_audit(path)
    return problems


def xml_link_audit(path):
    """python-pptx's click_action only sees the shape-level link. A tile carries a
    second one on its text run, and PowerPoint follows whichever the click lands
    on, so both have to be checked - at the XML level, where they both live."""
    import zipfile
    from lxml import etree
    NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    RID = f"{{{NS['r']}}}id"
    Z = zipfile.ZipFile(path)
    prels = {r.get("Id"): r.get("Target").split("/")[-1]
             for r in etree.fromstring(Z.read("ppt/_rels/presentation.xml.rels"))}
    order = [prels[e.get(RID)] for e in
             etree.fromstring(Z.read("ppt/presentation.xml")).iter(f"{{{NS['p']}}}sldId")]
    at = {f: i + 1 for i, f in enumerate(order)}
    out, seen = [], {}
    for n, f in enumerate(order, 1):
        rels = {r.get("Id"): r.get("Target") for r in
                etree.fromstring(Z.read(f"ppt/slides/_rels/{f}.rels"))}
        tree = etree.fromstring(Z.read(f"ppt/slides/{f}"))
        for sp in tree.iter(f"{{{NS['p']}}}sp"):
            txt = ("".join(x.text or "" for x in sp.iter(f"{{{NS['a']}}}t"))).strip()
            levels = {"shape": sp.find(f".//{{{NS['p']}}}nvSpPr"),
                      "text": sp.find(f".//{{{NS['p']}}}txBody")}
            dests = {}
            for lvl, root in levels.items():
                if root is None:
                    continue
                for hl in root.iter(f"{{{NS['a']}}}hlinkClick"):
                    rid = hl.get(RID)
                    if not rid:
                        out.append(f"slide {n} [{lvl}] '{txt[:24]}': link with no r:id"); continue
                    tgt = rels.get(rid)
                    if tgt is None:
                        out.append(f"slide {n} [{lvl}] '{txt[:24]}': r:id {rid} missing from rels"); continue
                    d = at.get(tgt.split("/")[-1])
                    if d is None:
                        out.append(f"slide {n} [{lvl}] '{txt[:24]}': target {tgt} is not a slide"); continue
                    if "hlinksldjump" not in (hl.get("action") or ""):
                        out.append(f"slide {n} [{lvl}] '{txt[:24]}': not a slide-jump action")
                    dests.setdefault(lvl, set()).add(d)
            if len(dests) == 2 and dests["shape"] != dests["text"]:
                out.append(f"slide {n} '{txt[:24]}': shape link -> {dests['shape']}, "
                           f"text link -> {dests['text']}")
            # a filled shape is clickable all over; a no-fill one only on its glyphs
            if dests and not txt.isdigit():
                spPr = sp.find(f".//{{{NS['p']}}}spPr")
                if spPr is not None and spPr.find(f"{{{NS['a']}}}noFill") is not None \
                        and "shape" in dests and "text" not in dests:
                    out.append(f"slide {n} '{txt[:24]}': button has no fill, so only the "
                               f"letters would be clickable")
    return out

DEFAULT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "materials", "jeopardy", "Jeopardy_Interesting_People_A2_fixed.pptx")

if __name__ == "__main__":
    probs = audit(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
    print(f"\n=== {len(probs)} problems ===")
    for x in probs: print("  -", x)
