# -*- coding: utf-8 -*-
"""Re-skin the Jeopardy board with the Great Britons artwork.

The artwork is used as the slide background exactly as drawn - the values,
category headers and title are painted into it. The forty tiles stay as shapes
on top, carrying the hyperlinks they already had, but become invisible hotspots
sitting precisely over the painted plaques.

A hotspot must have a FILL to be clickable: a shape set to "no fill" is
hit-tested only on its glyphs, so the tiles get a solid fill at 100 %
transparency instead, which covers the whole rectangle.
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    ROOT, "materials", "jeopardy", "Jeopardy_Interesting_People_A2_fixed.pptx")
BG = sys.argv[2]
GRID = json.load(open(sys.argv[3]))
OUT = sys.argv[4]

prs = Presentation(DECK)
slide = prs.slides[1]

# tiles are the shapes whose text is the point value; keep them, drop the rest
tiles, junk = [], []
for sh in slide.shapes:
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
    (tiles if txt.isdigit() else junk).append(sh)
assert len(tiles) == 40, f"expected 40 tiles, found {len(tiles)}"
print(f"плиток сохраняю: {len(tiles)}   прочих фигур удаляю: {len(junk)}")

# order tiles by their position, the same order the grid is generated in
def key(sh):
    return (round(sh.left / 100000), round(sh.top / 100000))
tiles.sort(key=lambda s: (round(s.left / 91440), round(s.top / 91440)))

for sh in junk:
    sh._element.getparent().remove(sh._element)

pic = slide.shapes.add_picture(BG, 0, 0, width=prs.slide_width, height=prs.slide_height)
spTree = slide.shapes._spTree
spTree.remove(pic._element)
spTree.insert(2, pic._element)          # behind everything, after nvGrpSpPr + grpSpPr
print(f"фон вставлен: {os.path.basename(BG)}")

def hotspot(sh, g):
    sh.left, sh.top = Inches(g["left"]), Inches(g["top"])
    sh.width, sh.height = Inches(g["width"]), Inches(g["height"])
    tf = sh.text_frame
    tf.clear()
    tf.paragraphs[0].text = ""
    spPr = sh._element.spPr
    for tag in ("solidFill", "noFill", "gradFill", "blipFill", "pattFill", "ln"):
        for el in spPr.findall(f"{{{A}}}{tag}"):
            spPr.remove(el)
    fill = etree.SubElement(spPr, f"{{{A}}}solidFill")
    clr = etree.SubElement(fill, f"{{{A}}}srgbClr"); clr.set("val", "FFFFFF")
    etree.SubElement(clr, f"{{{A}}}alpha").set("val", "0")     # invisible, still clickable
    ln = etree.SubElement(spPr, f"{{{A}}}ln")
    etree.SubElement(ln, f"{{{A}}}noFill")

for sh, g in zip(tiles, GRID):
    hotspot(sh, g)
print("плитки переставлены на нарисованные плашки и сделаны прозрачными")

prs.save(OUT)
print(f"записано: {OUT}")
