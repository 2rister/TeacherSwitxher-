# -*- coding: utf-8 -*-
"""Editable Word versions of the pack:  python3 build/make_docx.py

These are for editing, not for the poster look — Word cannot reproduce the
parchment, the engraved plates or the numbered colour bars. Layout, wording and
exercise order match the PDFs exactly, so a teacher can retype a question here
and print from Word, or paste changes back into build/content.py.
"""
import html as _html
import os, re, sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "materials", "docx")
IMG = os.path.join(ROOT, "materials", "img")
sys.path.insert(0, HERE)

import content as C
import grammar as G
import monday as M
import build as B

RED = RGBColor(0x9D, 0x22, 0x26)
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GREY = RGBColor(0x5A, 0x4A, 0x38)
BODY_FONT = "Cambria"


def txt(s):
    """content.py stores HTML; Word wants plain text."""
    return re.sub(r"\s+", " ", _html.unescape(re.sub(r"<[^>]+>", "", str(s)))).strip()


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    set_page(doc)
    return doc


def set_page(doc):
    """python-docx's default template is US Letter; the whole pack is A4, and a
    Letter page only shows itself when the file is actually opened."""
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(1.6)
        sec.left_margin = sec.right_margin = Cm(1.8)


def para(doc, text="", size=11, bold=False, italic=False, color=None,
         align=None, space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return p


def task_head(doc, no, title, hint):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"TASK {no}   {title.upper()}")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = NAVY
    if hint:
        h = p.add_run(f"   {hint}")
        h.italic = True
        h.font.size = Pt(9.5)
        h.font.color.rgb = GREY
    return p


def grid(doc, rows, cols, widths=None):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        t.autofit = False
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
    return t


def cell_text(cell, text, size=10.5, bold=False, italic=False, color=None, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color


L = B.L   # answer-letter helper, shared with the PDF build


def person_doc(d):
    doc = new_doc()

    para(doc, "INTERESTING PEOPLE  ·  READING & FACTS  ·  LEVEL A2",
         size=9, bold=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    head = grid(doc, 1, 2, widths=[4.2, 12.6])
    head.style = "Normal Table"
    pic = head.cell(0, 0).paragraphs[0]
    pic.add_run().add_picture(os.path.join(IMG, d["portrait"]), width=Cm(3.9))

    box = head.cell(0, 1)
    cell_text(box, txt(d["name"]).upper(), size=22, bold=True, color=NAVY)
    for line, kw in ((txt(d["dates"]), dict(size=11, bold=True, color=RED)),
                     (txt(d["strap"]), dict(size=10.5, color=GREY)),
                     (txt(d["tagline"]), dict(size=11.5, italic=True))):
        p = box.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        r.font.size = Pt(kw["size"])
        r.bold = kw.get("bold", False)
        r.italic = kw.get("italic", False)
        if kw.get("color") is not None:
            r.font.color.rgb = kw["color"]

    para(doc, "READ THE TEXT", size=13, bold=True, color=NAVY, space_before=12, space_after=2)
    para(doc, "Paragraph numbers help you find the answers.",
         size=9.5, italic=True, color=GREY, space_after=6)
    for i, t in enumerate(d["text"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        n = p.add_run(f"{i}  ")
        n.bold = True
        n.font.size = Pt(9)
        n.font.color.rgb = RED
        p.add_run(txt(t)).font.size = Pt(11)

    para(doc, f'"{txt(d["pull"])}"', size=12, italic=True, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=10)

    # ---- glossary
    gl = C.GLOSSARY[d["slug"]]
    para(doc, "WORDS TO KNOW", size=12, bold=True, color=NAVY, space_before=6, space_after=4)
    t = grid(doc, (len(gl) + 1) // 2, 2, widths=[8.4, 8.4])
    for i, (w, expl) in enumerate(gl):
        c = t.cell(i // 2, i % 2)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(txt(w))
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RED
        p.add_run(" – " + txt(expl)).font.size = Pt(10)

    doc.add_section(WD_SECTION.NEW_PAGE)

    # ---- 1  true / false / not given
    task_head(doc, 1, "True, False or Not Given", "Circle T, F or NG.")
    t = grid(doc, len(d["tf"]), 3, widths=[1.0, 13.2, 2.6])
    for i, (s, _) in enumerate(d["tf"]):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(s))
        cell_text(t.cell(i, 2), "T   F   NG", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ---- 2  gap fill
    task_head(doc, 2, "Fill in the gaps", "Use the words in the box.")
    bank = doc.add_table(rows=1, cols=1)
    bank.style = "Table Grid"
    shade(bank.cell(0, 0), "EFEFEF")
    cell_text(bank.cell(0, 0), "   ·   ".join(txt(w) for w in d["bank"]),
              size=10.5, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (s, _) in enumerate(d["gaps"], 1):
        p = doc.add_paragraph(style="List Number" if False else None)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Cm(0.7)
        p.add_run(f"{i}.  ").bold = True
        p.add_run(txt(s).replace("____________", "_" * 16)).font.size = Pt(11)

    doc.add_section(WD_SECTION.NEW_PAGE)

    # ---- 3  open questions
    task_head(doc, 3, "Answer the questions", "Write full sentences.")
    for i, (q, _) in enumerate(d["qs"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.7)
        p.add_run(f"{i}.  ").bold = True
        p.add_run(txt(q)).font.size = Pt(11)
        for _ in range(2):
            r = doc.add_paragraph()
            r.paragraph_format.space_after = Pt(6)
            r.paragraph_format.left_indent = Cm(1.2)
            r.add_run("." * 96).font.color.rgb = GREY

    # ---- 4  matching
    # _match_key[i] is the display slot holding stem i's half; invert it to redraw
    # the shuffled column exactly as the PDF prints it.
    rhs = [None] * len(d["match"])
    for i, slot in enumerate(d["_match_key"]):
        rhs[slot] = d["match"][i][1]
    task_head(doc, 4, "Match the two halves", "Write a, b, c, d or e.")
    t = grid(doc, len(d["match"]), 4, widths=[1.0, 1.3, 7.1, 7.4])
    for i, ((stem, _), half) in enumerate(zip(d["match"], rhs)):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), "", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 2), txt(stem))
        cell_text(t.cell(i, 3), f"{L(i)})  " + txt(half))

    # ---- 5  jeopardy call-back
    task_head(doc, 5, "Back to the board", "From the game – texts face down!")
    t = grid(doc, len(d["jeopardy"]), 2, widths=[2.0, 14.8])
    for i, (pts, q, _) in enumerate(d["jeopardy"]):
        shade(t.cell(i, 0), "1B2A4A")
        cell_text(t.cell(i, 0), str(pts), bold=True,
                  color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(q))

    return doc


def mixed_doc():
    doc = new_doc()
    para(doc, "THE MIXED ROUND  ·  NEWTON  ·  LENNON  ·  DIANA",
         size=13, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    task_head(doc, 1, "Who is it?", "N = Newton  ·  L = Lennon  ·  D = Diana")
    t = grid(doc, len(C.WHO), 3, widths=[1.0, 13.2, 2.6])
    for i, (q, _) in enumerate(C.WHO):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(q))
        cell_text(t.cell(i, 2), "N   L   D", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    task_head(doc, 2, "Put them in order", "1 = oldest, 5 = youngest. Use the year of birth.")
    t = grid(doc, len(C.TIMELINE), 2, widths=[1.6, 15.2])
    for i, (n, y) in enumerate(C.TIMELINE):
        cell_text(t.cell(i, 0), "", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), f"{txt(n)}  ({y})")

    task_head(doc, 3, "What do these numbers mean?", "Write a–h next to each number.")
    nums_rhs = [None] * len(C.NUMBERS)
    for i, slot in enumerate(C._nums_key):
        nums_rhs[slot] = C.NUMBERS[i][1]
    t = grid(doc, len(C.NUMBERS), 4, widths=[1.0, 1.3, 3.4, 11.1])
    for i, ((num, _), fact) in enumerate(zip(C.NUMBERS, nums_rhs)):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), "", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 2), txt(num), bold=True)
        cell_text(t.cell(i, 3), f"{L(i)})  " + txt(fact))

    doc.add_section(WD_SECTION.NEW_PAGE)

    task_head(doc, 4, "Final Jeopardy · Speaking", "Work in pairs, then tell the class.")
    para(doc, txt(C.DEBATE["question"]), size=12, bold=True, space_after=8)
    t = grid(doc, 2, len(C.DEBATE["steps"]), widths=[4.2] * len(C.DEBATE["steps"]))
    for i, (step, say) in enumerate(C.DEBATE["steps"]):
        shade(t.cell(0, i), "9D2226" if i % 2 == 0 else "1B2A4A")
        cell_text(t.cell(0, i), step, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(1, i), txt(say), italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    task_head(doc, 5, "Useful phrases", "Take one phrase from each box.")
    t = grid(doc, (len(C.DEBATE["phrases"]) + 1) // 2, 2, widths=[8.4, 8.4])
    for i, (title, phrases) in enumerate(C.DEBATE["phrases"]):
        c = t.cell(i // 2, i % 2)
        cell_text(c, txt(title), size=10, bold=True, color=RED)
        for ph in phrases:
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.add_run("• " + txt(ph)).italic = True
            p.runs[0].font.size = Pt(10)
    return doc


def key_doc():
    doc = new_doc()
    para(doc, "ANSWER KEY  ·  FOR THE TEACHER", size=14, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    for d in C.PEOPLE:
        para(doc, f"SHEET {d['num']}  ·  {txt(d['name']).upper()}",
             size=12, bold=True, color=NAVY, space_before=10, space_after=4)
        para(doc, "1  T/F/NG:  " + "   ".join(f"{i}{a}" for i, (_, a) in enumerate(d["tf"], 1)),
             size=10.5)
        para(doc, "2  Gaps:  " + "   ".join(f"{i} {txt(a)}" for i, (_, a) in enumerate(d["gaps"], 1)),
             size=10.5)
        para(doc, "3  Questions:", size=10.5, bold=True, space_after=1)
        for i, (_, a) in enumerate(d["qs"], 1):
            para(doc, f"     {i}. {txt(a)}", size=10.5, space_after=1)
        para(doc, "4  Matching:  " + "   ".join(f"{i} {L(k)}"
                                                for i, k in enumerate(d["_match_key"], 1)),
             size=10.5, space_before=3)
        para(doc, "5  Back to the board:", size=10.5, bold=True, space_after=1)
        for pts, _, a in d["jeopardy"]:
            para(doc, f"     {pts} pts – {txt(a)}", size=10.5, space_after=1)

    para(doc, "SHEET IV  ·  THE MIXED ROUND", size=12, bold=True, color=NAVY,
         space_before=12, space_after=4)
    para(doc, "1  Who is it?  " + "   ".join(f"{i} {a[0]}" for i, (_, a) in enumerate(C.WHO, 1)),
         size=10.5)
    para(doc, "2  Order:  " + "  →  ".join(txt(n) for n in C.TIMELINE_KEY), size=10.5)
    rank = {n: i for i, n in enumerate(C.TIMELINE_KEY, 1)}
    para(doc, "     Numbers to write in the boxes, top to bottom as printed:  "
              + "  ".join(str(rank[n]) for n, _ in C.TIMELINE), size=10, color=GREY)
    para(doc, "3  Numbers:  " + "   ".join(f"{i} {L(k)}"
                                           for i, k in enumerate(C._nums_key, 1)), size=10.5)
    para(doc, "4  Final Jeopardy:  No single correct answer. Accept any person if the student "
              "gives a reason and an example from the text.", size=10.5)

    doc.add_section(WD_SECTION.NEW_PAGE)
    para(doc, "TEACHER'S NOTES  ·  BEFORE YOU PRINT", size=14, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    for title, body in C.TEACHER_NOTES:
        para(doc, txt(title), size=11.5, bold=True, color=RED, space_before=8, space_after=2)
        para(doc, txt(body), size=10.5)

    para(doc, "IMAGE CREDITS & LICENCES", size=11.5, bold=True, color=RED,
         space_before=12, space_after=2)
    for line in (
        "Isaac Newton – portrait by Godfrey Kneller, 1689. Public domain. Via Wikimedia Commons.",
        "John Lennon – photograph by Joost Evers / Anefo, 1969. CC0 1.0. Via Wikimedia Commons.",
        "Princess Diana – photograph by John Mathew Smith, 1997. CC BY-SA 2.0; this pack "
        "reproduces it in duotone as a derivative work under the same licence. Via Wikimedia Commons.",
        "Type: Cinzel, Playfair Display, EB Garamond – SIL Open Font License 1.1.",
    ):
        para(doc, line, size=9.5, space_after=2)
    return doc


def grammar_doc(d, mixed=False):
    doc = new_doc()
    para(doc, "PAST SIMPLE  ·  GRAMMAR  ·  LEVEL A2", size=9, bold=True, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    para(doc, txt(d["name"]).upper(), size=18, bold=True, color=NAVY, space_after=1)
    para(doc, txt(d["dates"]), size=10.5, bold=True, color=RED, space_after=10)

    task_head(doc, 1, "Choose the correct answer", "Circle a, b or c.")
    for i, (sentence, opts, _) in enumerate(d["mc"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Cm(0.7)
        p.add_run(f"{i}.  ").bold = True
        p.add_run(txt(sentence).replace("____", "_" * 8)).font.size = Pt(11)
        for j, o in enumerate(opts):
            lab = p.add_run(f"    {chr(97 + j)}) ")
            lab.bold = True
            lab.font.size = Pt(10)
            lab.font.color.rgb = RED
            p.add_run(txt(o)).font.size = Pt(10.5)

    nxt = 2
    if mixed:
        task_head(doc, 2, "Make the question", "Put the words in the right order.")
        for i, (scram, _) in enumerate(G.WORD_ORDER, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.7)
            p.add_run(f"{i}.  ").bold = True
            r = p.add_run(txt(scram)); r.italic = True; r.font.size = Pt(11)
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(6)
            q.paragraph_format.left_indent = Cm(1.2)
            q.add_run("." * 96).font.color.rgb = GREY
        nxt = 3

    task_head(doc, nxt, "Put the verb into the Past Simple", "Use the verb in brackets.")
    for i, (sentence, _) in enumerate(d["bracket"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Cm(0.7)
        p.add_run(f"{i}.  ").bold = True
        p.add_run(txt(sentence).replace("__________", "_" * 18)).font.size = Pt(11)

    if mixed:
        task_head(doc, 4, "Irregular verbs", "Write the Past Simple form.")
        t = grid(doc, (len(G.IRREGULAR) + 1) // 2, 4, widths=[3.6, 4.8, 3.6, 4.8])
        for i, (inf, _) in enumerate(G.IRREGULAR):
            r, c = i // 2, (i % 2) * 2
            cell_text(t.cell(r, c), txt(inf), bold=True, color=RED)
            cell_text(t.cell(r, c + 1), "")
    return doc


def grammar_key_doc():
    doc = new_doc()
    para(doc, "GRAMMAR ANSWER KEY  ·  FOR THE TEACHER", size=14, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    for d in G.GRAMMAR_PEOPLE + [G.MIXED]:
        para(doc, f"SHEET {d['num']}  ·  {txt(d['name']).upper()}",
             size=12, bold=True, color=NAVY, space_before=10, space_after=4)
        para(doc, "1  Choose:  " + "   ".join(f"{i}{chr(97 + k)}"
                                              for i, (_, _, k) in enumerate(d["mc"], 1)), size=10.5)
        n = 2
        if d is G.MIXED:
            para(doc, "2  Make the question:", size=10.5, bold=True, space_after=1)
            for i, (_, a) in enumerate(G.WORD_ORDER, 1):
                para(doc, f"     {i}. {txt(a)}", size=10.5, space_after=1)
            n = 3
        para(doc, f"{n}  Past Simple:", size=10.5, bold=True, space_before=3, space_after=1)
        for i, (_, a) in enumerate(d["bracket"], 1):
            val = txt(a) if isinstance(a, str) else "   ".join(
                f"({j}) {txt(x)}" for j, x in enumerate(a, 1))
            para(doc, f"     {i}. {val}", size=10.5, space_after=1)

    para(doc, f"SHEET {G.MIXED['num']}  ·  TASK 4  ·  IRREGULAR VERBS", size=12, bold=True, color=NAVY,
         space_before=12, space_after=4)
    para(doc, "   ".join(f"{a} \u2192 {b}" for a, b in G.IRREGULAR), size=10.5)
    para(doc, "NOTES", size=12, bold=True, color=RED, space_before=12, space_after=3)
    para(doc, txt(G.GRAMMAR_NOTE), size=10.5)
    para(doc, "Both the full and the contracted form are correct: did not like and didn't like "
              "are equally right, and so are was not and wasn't. In task 'Make the question' and "
              "in question items, the capital letter is part of the answer.",
         size=10.5, space_before=6)
    return doc



# ====================================================== Monday 14.09 lesson kit
def kit_title(doc, title, sub):
    para(doc, "MONDAY 14 SEPTEMBER 2026  ·  ELIZABETH I  ·  WHAT MAKES A GREAT PERSON?",
         size=9, bold=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    para(doc, title.upper(), size=20, bold=True, color=NAVY, space_after=2)
    para(doc, sub.upper(), size=9.5, bold=True, color=RED, space_after=10)


def section_title(doc, t):
    para(doc, t.upper(), size=12, bold=True, color=NAVY, space_before=12, space_after=4)


def note_box(doc, title, body):
    t = grid(doc, 1, 1)
    shade(t.cell(0, 0), "F3F1EA")
    c = t.cell(0, 0)
    cell_text(c, txt(title), size=10.5, bold=True, color=RED)
    q = c.add_paragraph()
    q.paragraph_format.space_after = Pt(3)
    q.add_run(txt(body)).font.size = Pt(10)


def stage_grid(doc, rows):
    t = grid(doc, len(rows), 4, widths=[2.0, 1.3, 11.4, 2.1])
    for i, (clock, mins, what, how, ix) in enumerate(rows):
        cell_text(t.cell(i, 0), clock, size=10, bold=True, color=RED)
        cell_text(t.cell(i, 1), f"{mins}'", size=9, color=GREY)
        c = t.cell(i, 2)
        cell_text(c, txt(what), size=10.5, bold=True, color=NAVY)
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.add_run(txt(B.subst(how))).font.size = Pt(9.5)
        cell_text(t.cell(i, 3), ix.replace("&#8594;", "->"), size=8, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)


def monday_plan_doc():
    doc = new_doc()
    kit_title(doc, "Lesson plan", "Monday 14.09.2026 · two and a half sessions · level A2")

    section_title(doc, "The day")
    t = grid(doc, len(M.TIMETABLE), 3, widths=[4.2, 10.4, 2.2])
    for i, (a, b, c) in enumerate(M.TIMETABLE):
        cell_text(t.cell(i, 0), txt(a), bold=True, color=RED)
        cell_text(t.cell(i, 1), txt(b))
        cell_text(t.cell(i, 2), txt(c), size=9, color=GREY,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)

    section_title(doc, "By the end of the day students can")
    for a in M.AIMS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(txt(a)).font.size = Pt(10.5)

    section_title(doc, "On the table before you start")
    t = grid(doc, len(M.MATERIALS), 3, widths=[4.6, 10.0, 2.2])
    for i, (a, b, c) in enumerate(M.MATERIALS):
        cell_text(t.cell(i, 0), txt(B.subst(a)), bold=True, color=NAVY, size=10)
        cell_text(t.cell(i, 1), txt(B.subst(b)), size=10)
        cell_text(t.cell(i, 2), txt(c), size=9, color=GREY,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_section(WD_SECTION.NEW_PAGE)
    section_title(doc, "Session 1 · 11.00 – 11.45 · revision and Past Simple")
    stage_grid(doc, M.STAGE1)
    section_title(doc, "Session 2 · 12.00 – 13.35 · preparing the presentation")
    stage_grid(doc, M.STAGE2)

    doc.add_section(WD_SECTION.NEW_PAGE)
    section_title(doc, "Session 3 · 14.05 – 15.40 · plenary, presentations, awards")
    stage_grid(doc, M.STAGE3)
    section_title(doc, "Notes for the teacher")
    for t_, b_ in M.PLAN_NOTES:
        note_box(doc, t_, b_)
    return doc


def monday_listening_doc():
    doc = new_doc()
    kit_title(doc, "The Queen Who Said No",
              "Listening · A2 · you will hear the talk twice")

    rhs = [M.L_WORDS[M._lwords_key.index(i)][1] for i in range(len(M.L_WORDS))]
    task_head(doc, 1, "Before you listen", "Match the word to its meaning. Write a–h.")
    t = grid(doc, len(M.L_WORDS), 4, widths=[1.0, 5.2, 1.0, 9.6])
    for i, ((w, _), r) in enumerate(zip(M.L_WORDS, rhs)):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(w), bold=True)
        cell_text(t.cell(i, 2), f"{L(i)})", align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
        cell_text(t.cell(i, 3), txt(r))

    task_head(doc, 2, "First listening", "Tick the FIVE things the talk speaks about.")
    t = grid(doc, (len(M.L_TICK) + 1) // 2, 4, widths=[1.2, 7.2, 1.2, 7.2])
    for i, (s_, _) in enumerate(M.L_TICK):
        r, c = i % ((len(M.L_TICK) + 1) // 2), (i // ((len(M.L_TICK) + 1) // 2)) * 2
        cell_text(t.cell(r, c), "[  ]", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(r, c + 1), txt(s_))

    task_head(doc, 3, "Second listening", "Circle T, F or NG.")
    t = grid(doc, len(M.L_TF), 3, widths=[1.0, 13.2, 2.6])
    for i, (s_, _) in enumerate(M.L_TF):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(s_))
        cell_text(t.cell(i, 2), "T   F   NG", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_section(WD_SECTION.NEW_PAGE)
    nrhs = [M.L_NUMBERS[M._lnums_key.index(i)][1] for i in range(len(M.L_NUMBERS))]
    task_head(doc, 4, "The numbers", "Write a–g next to each number.")
    t = grid(doc, len(M.L_NUMBERS), 4, widths=[1.0, 4.0, 1.0, 10.8])
    for i, ((n_, _), r) in enumerate(zip(M.L_NUMBERS, nrhs)):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(n_), bold=True)
        cell_text(t.cell(i, 2), f"{L(i)})", align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
        cell_text(t.cell(i, 3), txt(r))

    task_head(doc, 5, "After you listen", "Talk in your team, then write.")
    for i, q in enumerate(M.L_AFTER, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{i}.  ")
        r.bold = True
        r.font.color.rgb = RED
        p.add_run(txt(q)).font.size = Pt(10.5)
        for _ in range(2):
            para(doc, "_" * 92, size=10.5, color=GREY, space_after=3)
    return doc


def monday_script_doc():
    doc = new_doc()
    kit_title(doc, "Audio script", "Read aloud by the teacher · for the teacher only")
    para(doc, txt(M.SCRIPT_META), size=9.5, italic=True, color=GREY, space_after=8)
    for i, (t_, pause) in enumerate(M.SCRIPT, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        n = p.add_run(f"{i}  ")
        n.bold = True
        n.font.size = Pt(9)
        n.font.color.rgb = RED
        p.add_run(txt(t_)).font.size = Pt(12)
        if pause:
            r = p.add_run("  ||")
            r.bold = True
            r.font.color.rgb = RED

    doc.add_section(WD_SECTION.NEW_PAGE)
    section_title(doc, "Listening sheet · answer key")
    rows = [
        ("1  Before you listen",
         " · ".join(f"{i} {L(k)}" for i, k in enumerate(M._lwords_key, 1))),
        ("2  First listening",
         " · ".join(f"{i} {'tick' if t_ else '–'}"
                    for i, (_, t_) in enumerate(M.L_TICK, 1))),
        ("3  Second listening",
         " · ".join(f"{i} {a}" for i, (_, a) in enumerate(M.L_TF, 1))),
        ("4  The numbers",
         " · ".join(f"{i} {L(k)}" for i, k in enumerate(M._lnums_key, 1))),
        ("5  After you listen", "Open answers."),
    ]
    t = grid(doc, len(rows), 2, widths=[4.6, 12.2])
    for i, (a, b) in enumerate(rows):
        cell_text(t.cell(i, 0), a, size=10, bold=True, color=NAVY)
        cell_text(t.cell(i, 1), b, size=10)

    section_title(doc, "Notes")
    note_box(doc, "Why item 7 of task 3 is Not Given",
             "The talk says Shakespeare wrote his plays in her time. It never says she met "
             "him or wrote with him. Students who answer F are reading the world, not the "
             "text; students who answer T are reading a film.")
    note_box(doc, "How to read it",
             "First reading: normal speed, no stopping, students only tick task 2. Second "
             "reading: slow down a little and pause at ||. If the group is weak, read "
             "paragraph 7 a third time on its own.")
    note_box(doc, "Three more ways to use the same script",
             "Running dictation (cut into ten paragraphs, tape to the walls, one runs and "
             "dictates, one writes). Shadowing paragraph 7 line by line, then one volunteer "
             "says the quotation standing up. Retell relay: one sentence each in the Past "
             "Simple, wrong tense sends the chain back to the start.")
    note_box(doc, "Where the facts come from (checked 13.09.2026)", M.TEACHER_EVIDENCE)
    return doc


def monday_great_doc():
    doc = new_doc()
    kit_title(doc, "What makes a great person?",
              "Session 2 · build your criteria, then test them")

    task_head(doc, 1, "Your criteria", "Tick five alone. Then agree on three as a team.")
    n = (len(M.CRITERIA) + 1) // 2
    t = grid(doc, n, 4, widths=[1.2, 7.2, 1.2, 7.2])
    for i, c_ in enumerate(M.CRITERIA):
        r, c = i % n, (i // n) * 2
        cell_text(t.cell(r, c), "[  ]", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(r, c + 1), txt(c_), size=10)

    para(doc, "", space_after=4)
    t = grid(doc, 4, 2, widths=[5.8, 11.0])
    cell_text(t.cell(0, 0), "OUR TEAM'S THREE CRITERIA", size=9.5, bold=True, color=NAVY)
    cell_text(t.cell(0, 1), "WHY THIS ONE AND NOT ANOTHER?", size=9.5, bold=True, color=NAVY)
    for r in range(1, 4):
        for c in range(2):
            cell_text(t.cell(r, c), "")

    task_head(doc, 2, "Evidence hunt", "No year, no evidence.")
    t = grid(doc, 4, 4, widths=[0.9, 4.6, 8.7, 2.6])
    for j, h in enumerate(("#", "CRITERION", "WHAT ELIZABETH ACTUALLY DID", "YEAR")):
        cell_text(t.cell(0, j), h, size=9.5, bold=True, color=NAVY)
    for r in range(1, 4):
        cell_text(t.cell(r, 0), str(r), align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
        for c in range(1, 4):
            cell_text(t.cell(r, c), "\n")

    doc.add_section(WD_SECTION.NEW_PAGE)
    task_head(doc, 3, "The greatness test",
              "Four things she really did. Score each 1–3 against YOUR criteria.")
    t = grid(doc, len(M.GREATNESS_TEST), 3, widths=[0.9, 13.3, 2.6])
    for i, (what, so) in enumerate(M.GREATNESS_TEST):
        cell_text(t.cell(i, 0), f"{i + 1}.", align=WD_ALIGN_PARAGRAPH.CENTER)
        c = t.cell(i, 1)
        cell_text(c, txt(what), size=10)
        q = c.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        rr = q.add_run(txt(so))
        rr.italic = True
        rr.font.size = Pt(9.5)
        rr.font.color.rgb = GREY
        cell_text(t.cell(i, 2), "1   2   3", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    task_head(doc, 4, "Both sides",
              "A team that cannot argue the other side has not understood its own.")
    para(doc, f'Motion: "{txt(M.DEBATE_MOTION)}"  Write three arguments on each side.',
         size=10.5, space_after=4)
    t = grid(doc, 4, 2, widths=[8.4, 8.4])
    cell_text(t.cell(0, 0), "FOR – BECAUSE SHE …", size=9.5, bold=True, color=NAVY)
    cell_text(t.cell(0, 1), "AGAINST – BUT SHE …", size=9.5, bold=True, color=NAVY)
    for r in range(1, 4):
        for c in range(2):
            cell_text(t.cell(r, c), "\n")
    note_box(doc, "Where the hard facts in task 3 come from", M.TEACHER_EVIDENCE)
    return doc


def monday_stage_doc():
    doc = new_doc()
    kit_title(doc, "Ten minutes on stage",
              "Six roles · everybody speaks · nobody runs over")

    task_head(doc, 1, "Who does what", "Write a real name in the last column.")
    t = grid(doc, len(M.ROLES), 5, widths=[0.8, 3.4, 2.6, 6.4, 3.6])
    for i, (no, role, tm, what) in enumerate(M.ROLES):
        cell_text(t.cell(i, 0), no, bold=True, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(i, 1), txt(role), bold=True, color=NAVY, size=10)
        cell_text(t.cell(i, 2), txt(tm), size=9.5, color=RED)
        cell_text(t.cell(i, 3), txt(what), size=9.5)
        cell_text(t.cell(i, 4), "", size=9.5)

    task_head(doc, 2, "Rehearsal checklist", "One student holds this and ticks during the run.")
    n = (len(M.CHECKLIST) + 1) // 2
    t = grid(doc, n, 4, widths=[1.2, 7.2, 1.2, 7.2])
    for i, c_ in enumerate(M.CHECKLIST):
        r, c = i % n, (i // n) * 2
        cell_text(t.cell(r, c), "[  ]", align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(r, c + 1), txt(c_), size=10)

    task_head(doc, 3, "The fix list", "Three things. Three.")
    t = grid(doc, 4, 1, widths=[16.8])
    cell_text(t.cell(0, 0), "AFTER REHEARSAL 1 – THREE THINGS WE CHANGE (NOT FOUR)",
              size=9.5, bold=True, color=NAVY)
    for r in range(1, 4):
        cell_text(t.cell(r, 0), "\n")

    doc.add_section(WD_SECTION.NEW_PAGE)
    section_title(doc, "Sentence frames – steal these")
    for title, xs in M.FRAMES:
        para(doc, txt(title).upper(), size=10, bold=True, color=RED,
             space_before=8, space_after=2)
        for x in xs:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            p.add_run(txt(x)).font.size = Pt(10)
    note_box(doc, "The one rule about the Past Simple on stage",
             "Everything that happened before today is Past Simple: she was, she became, "
             "she did not marry, did she marry? Everything you think now is present: we "
             "think, this shows, our answer is. Mixing the two is the single mistake the "
             "room will hear.")
    note_box(doc, "What a ten-minute talk is not",
             "It is not a Wikipedia page read out loud. The story of the life is two minutes "
             "of the ten. The other eight are your answer to the question, and the evidence "
             "for it.")
    return doc


def monday_cards_doc():
    doc = new_doc()
    kit_title(doc, "Audience cards", "Cut into four · one card per team you watch")
    for card in range(4):
        t = grid(doc, len(M.FB_SCORES) + len(M.FB_OPEN) + 2, 2, widths=[11.4, 5.4])
        cell_text(t.cell(0, 0), "AUDIENCE CARD", size=11, bold=True, color=NAVY)
        cell_text(t.cell(0, 1), "", size=9)
        cell_text(t.cell(1, 0), "Team: ____________________", size=9.5, color=GREY)
        cell_text(t.cell(1, 1), "Their person: ______________", size=9.5, color=GREY)
        for i, sc in enumerate(M.FB_SCORES):
            cell_text(t.cell(2 + i, 0), txt(sc), size=9.5)
            cell_text(t.cell(2 + i, 1), "1   2   3   4   5", size=9.5,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        base = 2 + len(M.FB_SCORES)
        for i, (lab, _) in enumerate(M.FB_OPEN):
            c = t.cell(base + i, 0)
            cell_text(c, txt(lab), size=9.5, bold=True)
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(2)
            q.add_run("_" * 60).font.size = Pt(9.5)
            cell_text(t.cell(base + i, 1), "", size=9.5)
        para(doc, "— — — — — — — — — — — — —  cut here  — — — — — — — — — — — — —",
             size=8, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER,
             space_before=4, space_after=6)

    doc.add_section(WD_SECTION.NEW_PAGE)
    section_title(doc, "Your three votes – tear off and post at 15.20")
    t = grid(doc, 3, 3, widths=[5.6, 5.6, 5.6])
    for i, (title, why) in enumerate(M.VOTES):
        cell_text(t.cell(0, i), txt(title), size=11, bold=True, color=RED,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(1, i), txt(why), size=9, color=GREY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.cell(2, i), "______________", align=WD_ALIGN_PARAGRAPH.CENTER)
    return doc


def monday_certs_doc():
    doc = new_doc()
    for i, (title, why, grp) in enumerate(M.CERTS):
        if i:
            doc.add_section(WD_SECTION.NEW_PAGE)
        para(doc, "", space_after=30)
        para(doc, txt(title), size=22, bold=True, color=RED,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
        para(doc, txt(why), size=11, italic=True, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
        para(doc, "AWARDED TO", size=9.5, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        para(doc, "_" * 46, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        para(doc, txt(grp), size=11, bold=True, color=RED,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
        t = grid(doc, 1, 2, widths=[8.4, 8.4])
        t.style = "Normal Table"
        cell_text(t.cell(0, 0), "Teacher  ______________________", size=9, color=GREY)
        cell_text(t.cell(0, 1), "Date  ______________________", size=9, color=GREY,
                  align=WD_ALIGN_PARAGRAPH.RIGHT)
        para(doc, txt(M.CERT_FOOT), size=8.5, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16)
    return doc


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    B.write_html()          # fixes the matching and numbers keys, exactly as the PDFs use them
    print("\nWriting editable Word versions\n")
    # same running order as the PDFs, so the two sets never drift apart
    seq = ([(d["slug"], person_doc(d)) for d in C.PEOPLE]
           + [("mixed-round", mixed_doc()), ("answer-key", key_doc())]
           + [(f"grammar-{d['slug']}", grammar_doc(d)) for d in G.GRAMMAR_PEOPLE]
           + [("grammar-mixed", grammar_doc(G.MIXED, mixed=True)),
              ("grammar-answer-key", grammar_key_doc()),
              ("monday-lesson-plan", monday_plan_doc()),
              ("monday-listening", monday_listening_doc()),
              ("monday-audio-script", monday_script_doc()),
              ("monday-great-person", monday_great_doc()),
              ("monday-presentation-kit", monday_stage_doc()),
              ("monday-audience-cards", monday_cards_doc()),
              ("monday-certificates", monday_certs_doc())])
    docs = [(f"{i:02d}-{name}", doc) for i, (name, doc) in enumerate(seq, 1)]
    for name, doc in docs:
        path = os.path.join(OUT, name + ".docx")
        set_page(doc)                      # covers sections added after new_doc()
        doc.save(path)
        print(f"  docx  {name}.docx  ({os.path.getsize(path) // 1024} KB)")
    print("\nDone.")
